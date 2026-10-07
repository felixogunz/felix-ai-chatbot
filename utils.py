import os
import re
import uuid
from pathlib import Path

from werkzeug.utils import secure_filename



# CONVERSATION TITLE


def generate_title(message):
    """
    Generate a short title from the user's first message.
    """

    if not message:
        return "New Chat"

    message = str(message).strip()

    if not message:
        return "New Chat"

    message = re.sub(r"\s+", " ", message)

    if len(message) <= 50:
        return message

    return message[:50].rstrip() + "..."



# SAFE FILE NAME


def get_safe_filename(filename):
    """
    Convert a filename into a safe filename.
    """

    if not filename:
        return "uploaded_file"

    filename = os.path.basename(
        str(filename)
    )

    filename = secure_filename(
        filename
    )

    if not filename:
        return "uploaded_file"

    return filename



# UNIQUE FILE NAME


def get_unique_filename(filename):
    """
    Generate a unique filename.
    """

    safe_filename = get_safe_filename(
        filename
    )

    file_path = Path(
        safe_filename
    )

    name = file_path.stem
    extension = file_path.suffix

    unique_id = uuid.uuid4().hex[:8]

    return (
        f"{name}_{unique_id}{extension}"
    )



# FILE EXTENSION


def get_file_extension(filename):
    """
    Return a file extension in lowercase.
    """

    if not filename:
        return ""

    return Path(
        str(filename)
    ).suffix.lower()



# ALLOWED FILE


def allowed_file(
    filename,
    allowed_extensions=None
):
    """
    Check whether a file has an allowed extension.
    """

    if not filename:
        return False

    if allowed_extensions is None:

        allowed_extensions = {
            ".txt",
            ".pdf",
            ".doc",
            ".docx"
        }

    extension = get_file_extension(
        filename
    )

    return extension in allowed_extensions



# FILE TYPE


def get_file_type(filename):
    """
    Return the general type of a file.
    """

    extension = get_file_extension(
        filename
    )

    file_types = {
        ".txt": "text",
        ".pdf": "pdf",
        ".doc": "document",
        ".docx": "document",
        ".jpg": "image",
        ".jpeg": "image",
        ".png": "image",
        ".gif": "image",
        ".webp": "image",
        ".csv": "spreadsheet",
        ".xls": "spreadsheet",
        ".xlsx": "spreadsheet"
    }

    return file_types.get(
        extension,
        "file"
    )



# FILE SIZE


def format_file_size(size):
    """
    Convert bytes into a readable file size.
    """

    try:

        size = float(size)

    except (TypeError, ValueError):

        return "0 B"

    if size < 1024:

        return f"{size:.0f} B"

    if size < 1024 ** 2:

        return f"{size / 1024:.2f} KB"

    if size < 1024 ** 3:

        return (
            f"{size / (1024 ** 2):.2f} MB"
        )

    if size < 1024 ** 4:

        return (
            f"{size / (1024 ** 3):.2f} GB"
        )

    return (
        f"{size / (1024 ** 4):.2f} TB"
    )



# CLEAN TEXT


def clean_text(text):
    """
    Remove unnecessary whitespace.
    """

    if not text:
        return ""

    text = str(text)

    lines = []

    for line in text.splitlines():

        line = line.strip()

        if line:
            lines.append(line)

    return "\n".join(lines)



# TRUNCATE TEXT


def truncate_text(
    text,
    max_length=12000
):
    """
    Limit text before sending it to the AI.
    """

    if not text:
        return ""

    text = str(text)

    if len(text) <= max_length:
        return text

    return (
        text[:max_length].rstrip()
        + "\n\n[Document truncated]"
    )



# EXTRACT PDF TEXT


def extract_pdf_text(file_path):
    """
    Extract text from a PDF.

    The PDF library is imported only when
    this function is actually used.
    """

    try:

        from pypdf import PdfReader

    except ImportError:

        try:

            from pypdf import PdfReader
            
        except ImportError:

            raise ImportError(
                "PDF support is not installed. "
                "Run: pip install pypdf"
            )#

    try:

        reader = PdfReader(
            str(file_path)
        )

        extracted_text = []

        for page in reader.pages:

            try:

                page_text = page.extract_text()

            except Exception:

                page_text = None

            if page_text:

                extracted_text.append(
                    page_text
                )

        return clean_text(
            "\n".join(
                extracted_text
            )
        )

    except Exception as error:

        raise RuntimeError(
            f"Unable to read PDF file: {error}"
        )



# EXTRACT DOCX TEXT


def extract_docx_text(file_path):
    """
    Extract text from a DOCX file.

    python-docx is imported only when
    this function is actually used.
    """

    try:

        from docx import Document

    except ImportError:

        raise ImportError(
            "DOCX support is not installed. "
            "Run: pip install python-docx"
        )

    try:

        document = Document(
            str(file_path)
        )

        paragraphs = []

        
        # NORMAL PARAGRAPHS
        

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(
                    text
                )

        
        # TABLES
        

        for table in document.tables:

            for row in table.rows:

                row_text = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:

                        row_text.append(
                            cell_text
                        )

                if row_text:

                    paragraphs.append(
                        " | ".join(
                            row_text
                        )
                    )

        return clean_text(
            "\n".join(
                paragraphs
            )
        )

    except Exception as error:

        raise RuntimeError(
            f"Unable to read DOCX file: {error}"
        )



# EXTRACT TEXT FROM FILE


def extract_text_from_file(file_path):
    """
    Extract text from TXT, PDF or DOCX files.
    """

    if not file_path:

        raise ValueError(
            "No file path supplied."
        )

    file_path = Path(
        file_path
    )

    if not file_path.exists():

        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if not file_path.is_file():

        raise ValueError(
            f"Path is not a file: {file_path}"
        )

    extension = file_path.suffix.lower()

    
    # TXT
    

    if extension == ".txt":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as file:

                return clean_text(
                    file.read()
                )

        except OSError as error:

            raise RuntimeError(
                f"Unable to read text file: {error}"
            )

    
    # PDF
    

    if extension == ".pdf":

        return extract_pdf_text(
            file_path
        )

    
    # DOCX
    

    if extension == ".docx":

        return extract_docx_text(
            file_path
        )

    
    # UNSUPPORTED FILE
    

    raise ValueError(
        f"Unsupported file type: {extension}"
    )



# ENSURE DIRECTORY

def ensure_directory(directory):
    """
    Create a directory if it doesn't exist.
    """

    if not directory:

        return False

    try:

        Path(directory).mkdir(
            parents=True,
            exist_ok=True
        )

        return True

    except OSError:

        return False



# SAFE PATH


def get_safe_path(
    directory,
    filename
):
    """
    Create a safe path inside a directory.
    """

    if not directory:

        raise ValueError(
            "Directory is required."
        )

    safe_filename = get_safe_filename(
        filename
    )

    base_directory = Path(
        directory
    ).resolve()

    file_path = (
        base_directory / safe_filename
    ).resolve()

    try:

        file_path.relative_to(
            base_directory
        )

    except ValueError:

        raise ValueError(
            "Invalid file path."
        )

    return file_path



# FILE EXISTS


def file_exists(file_path):
    """
    Check whether a file exists.
    """

    if not file_path:

        return False

    return Path(
        file_path
    ).is_file()



# READ TEXT FILE


def read_text_file(file_path):
    """
    Read a text file.
    """

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()

    except OSError:

        return ""



# FILE INFORMATION


def get_file_info(file_path):
    """
    Return basic information about a file.
    """

    if not file_path:

        return None

    path = Path(
        file_path
    )

    if not path.exists():

        return None

    try:

        size = path.stat().st_size

    except OSError:

        size = 0

    return {
        "name": path.name,
        "extension": path.suffix.lower(),
        "size": size,
        "size_formatted": format_file_size(size),
        "type": get_file_type(
            path.name
        )
    }